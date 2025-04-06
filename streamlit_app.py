import streamlit as st
import requests
import spacy
import os
from langchain_google_genai import ChatGoogleGenerativeAI
from langchain.schema import HumanMessage
from fpdf import FPDF
from docx import Document

# Set up API keys
os.environ["GOOGLE_API_KEY"] = st.secrets["GOOGLE_API_KEY"]
os.environ["REPLICATE_API_TOKEN"] = st.secrets["REPLICATE_API_TOKEN"]

# Load SpaCy model once
@st.cache_resource
def load_spacy_model():
    return spacy.load("en_core_sci_md")

nlp = load_spacy_model()

# Set up LLM
llm = ChatGoogleGenerativeAI(model="gemini-1.5-flash")

# Patient database persistence
if "patient_db" not in st.session_state:
    st.session_state.patient_db = {}

patient_db = st.session_state.patient_db

st.title("🩺 AI-Powered Medical Documentation System")
st.write("Upload an audio file of a patient consultation to generate a structured medical report.")

st.subheader("Patient Information")

pid = st.text_input("Enter Patient ID (if existing) or leave blank to register a new patient")

if pid:
    if pid in patient_db:
        st.success(f"Patient {pid} found!")
        patient = patient_db[pid]
        st.write(f"**Name:** {patient['name']}")
        st.write(f"**Age/Sex:** {patient['age']}/{patient['sex']}")
    else:
        st.warning("Patient ID not found. Please register the patient below.")
        name = st.text_input("Patient Name")
        age = st.number_input("Age", min_value=0, max_value=120, step=1)
        sex = st.radio("Sex", ["Male", "Female", "Other"])
        if st.button("Register Patient"):
            patient_db[pid] = {"name": name, "age": age, "sex": sex}
            st.success(f"Patient {pid} registered!")

st.subheader("Upload Consultation Audio")
audio_file = st.file_uploader("Upload an audio file (.mp3, .wav)")

if audio_file is not None:
    st.audio(audio_file, format='audio/wav')

    # Here you would normally use a transcription API or model
    # But for now, simulate transcription
    st.info("Transcribing audio (mockup)...")
    transcription = "The patient complains of chest pain and shortness of breath."

    st.subheader("Extracted Medical Information")

    # Extract entities using SciSpaCy
    doc = nlp(transcription)
    entities = [(ent.text, ent.label_) for ent in doc.ents]

    if entities:
        st.write("**Identified Medical Terms:**")
        for text, label in entities:
            st.write(f"- {text} ({label})")
    else:
        st.write("No medical entities found.")

    # Generate a medical report
    st.subheader("Generated Medical Report")
    messages = [HumanMessage(content=f"Summarize the following consultation in a medical report format:\n{transcription}")]
    report = llm(messages).content
    st.write(report)

    # Allow Download
    if st.button("Download Report as PDF"):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        pdf.multi_cell(0, 10, report)
        pdf_output_path = "/tmp/medical_report.pdf"
        pdf.output(pdf_output_path)

        with open(pdf_output_path, "rb") as f:
            st.download_button("Download PDF", f, file_name="medical_report.pdf")

    if st.button("Download Report as DOCX"):
        docx = Document()
        docx.add_heading("Medical Report", 0)
        docx.add_paragraph(report)
        docx_output_path = "/tmp/medical_report.docx"
        docx.save(docx_output_path)

        with open(docx_output_path, "rb") as f:
            st.download_button("Download DOCX", f, file_name="medical_report.docx")
